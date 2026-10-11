"""Redacted files must not give the masked values back, by any route."""

import io
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont

from anonimizador import ocr
from anonimizador.align import MaskedWord
from anonimizador.extract import extract_lines
from anonimizador.redact import RedactionError, redact
from anonimizador.redact_verify import _check_pixels
from anonimizador.sources import load_source
from anonimizador.verify import find_values

OFFICIAL = Path(__file__).resolve().parents[2] / "raw_data" / \
    "document_to_anonymize.pdf"
SECRET = {"Maria", "Santos", "123456789", "maria@exemplo.pt"}
FONTS = [Path(r"C:\Windows\Fonts\arial.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/System/Library/Fonts/Supplemental/Arial.ttf")]
FONT = next((f for f in FONTS if f.exists()), None)
needs_ocr = pytest.mark.skipif(
    not ocr.available() or FONT is None,
    reason="Tesseract with Portuguese or a TrueType font is missing")


def secret_words(lines):
    """Mask every word that is (or starts with) a secret value."""
    return [MaskedWord(n, i, w) for n, line in enumerate(lines)
            for i, w in enumerate(line.split())
            if w.strip(",.;:()") in SECRET]


def document_image():
    font = ImageFont.truetype(str(FONT), 40)
    image = Image.new("RGB", (1500, 420), "white")
    draw = ImageDraw.Draw(image)
    for row, text in enumerate([
            "Relatório de Admissão",
            "Nome: Maria Santos",
            "NIF: 123456789",
            "Email: maria@exemplo.pt"]):
        draw.text((50, 40 + row * 90), text, fill="black", font=font)
    return image


# --- leak rule ---------------------------------------------------------------

def test_find_values_ignores_layout_and_plain_words():
    lines = ["Doente: Maria", "Santos, NIF 123 456", "Lisboa"]
    assert find_values(["Maria Santos", "Lisboa", "123456"], lines) == [
        "Maria Santos"]
    assert find_values(["maria@x.pt"], ["Email: MARIA@X.PT."]) == [
        "maria@x.pt"]


def test_unpainted_zone_fails_the_pixel_check():
    image = Image.new("RGB", (50, 50), "white")
    square = [(10, 10), (20, 10), (20, 20), (10, 20)]
    with pytest.raises(RedactionError):
        _check_pixels(image, [square])


# --- images ------------------------------------------------------------------

@needs_ocr
def test_photo_comes_back_black_boxed_without_metadata():
    exif = Image.Exif()
    exif[0x010E] = "Maria Santos 123456789"   # ImageDescription
    exif[0x013B] = "Maria Santos"             # Artist
    buffer = io.BytesIO()
    document_image().save(buffer, format="JPEG", exif=exif, quality=95)

    source = load_source(buffer.getvalue(), ".jpg")
    assert source.kind == "image"
    assert any("Maria Santos" in line for line in source.lines)
    masked = secret_words(source.lines)
    assert len(masked) >= 3

    result = redact(source, masked, "foto.jpg")
    assert result.filename == "foto_censurado.png"
    assert b"Maria" not in result.data and b"123456789" not in result.data
    out = Image.open(io.BytesIO(result.data))
    assert out.format == "PNG" and not out.info and not out.getexif()
    lines = ocr.line_texts(ocr.read_image(out.convert("RGB")))
    assert not any("Maria" in line or "1234" in line for line in lines)
    assert any("Relatório" in line for line in lines)


@needs_ocr
def test_unlocated_word_fails_closed():
    buffer = io.BytesIO()
    document_image().save(buffer, format="PNG")
    source = load_source(buffer.getvalue(), ".png")
    with pytest.raises(RedactionError):
        redact(source, [MaskedWord(99, 0, "x")], "foto.png")


@needs_ocr
def test_rotated_photo_is_read_upright():
    exif = Image.Exif()
    exif[0x0112] = 6  # rotate 90 degrees when shown
    buffer = io.BytesIO()
    document_image().rotate(90, expand=True).save(buffer, format="JPEG",
                                                  exif=exif)
    lines = load_source(buffer.getvalue(), ".jpg").lines
    assert any("Maria Santos" in line for line in lines)


# --- PDFs --------------------------------------------------------------------

def text_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Nome: Maria Santos, NIF 123456789")
    page.insert_text((72, 100), "Diagnóstico: hipertensão")
    # The same name again in white: invisible, but extractable.
    page.insert_text((72, 128), "Maria Santos", color=(1, 1, 1))
    doc.set_metadata({"author": "Maria Santos", "title": "123456789"})
    return doc.tobytes()


def test_official_pdf_is_read_as_before():
    if not OFFICIAL.exists():
        pytest.skip("official PDF missing")
    source = load_source(OFFICIAL.read_bytes(), ".pdf")
    assert source.kind == "pdf"
    assert source.lines == extract_lines(OFFICIAL)
    words = sum(len(line.split()) for line in source.lines)
    assert len(source.locations) == words


def test_text_pdf_redaction_removes_text_hidden_text_and_metadata():
    source = load_source(text_pdf(), ".pdf")
    masked = secret_words(source.lines)
    result = redact(source, masked, "relatorio.pdf")
    with pymupdf.open(stream=result.data, filetype="pdf") as doc:
        text = "".join(page.get_text() for page in doc)
        assert "Maria" not in text and "123456789" not in text
        assert "hipertensão" in text
        assert not doc.metadata["author"] and not doc.metadata["title"]
    assert b"123456789" not in result.data and b"Maria" not in result.data


@needs_ocr
def test_scanned_pdf_is_rebuilt_from_redacted_images():
    buffer = io.BytesIO()
    document_image().save(buffer, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page(width=600, height=170)
    page.insert_image(page.rect, stream=buffer.getvalue())
    source = load_source(doc.tobytes(), ".pdf")
    assert source.kind == "scanned"
    result = redact(source, secret_words(source.lines), "scan.pdf")
    with pymupdf.open(stream=result.data, filetype="pdf") as out:
        assert not "".join(p.get_text() for p in out).strip()


# --- DOCX --------------------------------------------------------------------

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def make_docx() -> bytes:
    document = Document()
    document.core_properties.author = "Maria Santos"
    document.sections[0].header.paragraphs[0].text = "Processo de Maria Santos"
    paragraph = document.add_paragraph("Nome: Ma")
    bold = paragraph.add_run("ria Santos")
    bold.bold = True
    bold.font.size = Pt(14)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "NIF"
    table.cell(0, 1).text = "123456789"
    note = document.add_paragraph("Contacto: maria@exemplo.pt")
    document.add_comment(note.runs, text="Ligar à Maria Santos",
                         author="Maria Santos")
    deleted = document.add_paragraph("Notas: ")
    deleted._p.append(parse_xml(
        f'<w:del {W} w:id="1" w:author="Maria Santos" '
        f'w:date="2025-01-01T00:00:00Z"><w:r>'
        f'<w:delText>Maria Santos 123456789</w:delText></w:r></w:del>'))
    deleted.add_run("sem alterações")
    image = io.BytesIO()
    Image.new("RGB", (20, 20), "red").save(image, format="PNG")
    document.add_picture(image)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def all_xml_text(data: bytes) -> str:
    archive = zipfile.ZipFile(io.BytesIO(data))
    return "".join(archive.read(n).decode("utf-8", "replace")
                   for n in archive.namelist()
                   if n.endswith((".xml", ".rels")))


def test_docx_keeps_formatting_and_leaves_nothing_behind():
    source = load_source(make_docx(), ".docx")
    assert "Processo de Maria Santos" in source.lines
    assert "Nome: Maria Santos" in source.lines
    assert "Notas: sem alterações" in source.lines
    result = redact(source, secret_words(source.lines), "ficha.docx")

    xml = all_xml_text(result.data)
    assert "Maria" not in xml and "123456789" not in xml
    assert "maria@exemplo.pt" not in xml
    names = zipfile.ZipFile(io.BytesIO(result.data)).namelist()
    assert not any(n.startswith("word/media/") or "comments" in n
                   for n in names)

    out = load_source(result.data, ".docx")
    assert "Nome: * *" in out.lines and "Processo de * *" in out.lines
    runs = Document(io.BytesIO(result.data)).paragraphs[0].runs
    assert runs[1].bold and runs[1].font.size == Pt(14)


def names_only(prompt, model=None, json_output=False, **kwargs):
    """A model that masks the names and nothing else."""
    import json
    import re
    lines = re.findall(r"^L\d+: (.*)$", prompt, re.M)
    masked = [re.sub(r"Maria|Santos", "*", line) for line in lines]
    return json.dumps({"masked_lines": masked, "entities": []})


@needs_ocr
def test_misread_email_is_still_removed_from_the_photo():
    """OCR reads "maria@exemplo.pt" as "mariaQgexemplo.pt": neither the
    model nor the patterns see an email, so the low-confidence rule must."""
    from anonimizador.pipeline import anonymize_source

    buffer = io.BytesIO()
    document_image().save(buffer, format="JPEG", quality=90)
    source = load_source(buffer.getvalue(), ".jpg")
    result = anonymize_source(source, generate_fn=names_only)
    assert result.lines[3] == "Email: *"
    redacted = redact(source, result.masked, "foto.jpg")
    out = Image.open(io.BytesIO(redacted.data)).convert("RGB")
    text = " ".join(ocr.line_texts(ocr.read_image(out)))
    assert "exemplo" not in text and "Maria" not in text


@needs_ocr
def test_a_line_ocr_missed_is_detected_and_blocks_the_file():
    image = document_image()
    lines = ocr.read_image(image)
    assert ocr.unread_lines(image, lines) == 0
    missing_nif = [line for line in lines
                   if not any(w.text == "NIF:" for w in line)]
    assert ocr.unread_lines(image, missing_nif) >= 1

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    source = load_source(buffer.getvalue(), ".png")
    source.unread = 1
    with pytest.raises(RedactionError, match="não foram lidas"):
        redact(source, secret_words(source.lines), "foto.png")


def first_words(lines):
    return [line.split()[0] for line in ocr.line_texts(lines)]


@needs_ocr
def test_a_line_ocr_skipped_is_read_again_and_put_back_in_order():
    image = document_image()
    lines = ocr.read_image(image)
    missing_nif = [line for line in lines
                   if not any(w.text == "NIF:" for w in line)]
    merged, still_unread = ocr.recover_unread(image, missing_nif)
    assert still_unread == []
    # OCR may misread a digit when reading the line alone ("123456/89").
    assert first_words(merged) == first_words(lines)
    assert ocr.unread_lines(image, merged) == 0


def reads_back(lines):
    """A stand-in for reading a strip again: the words of ``lines`` whose
    boxes cross the strip, as a perfect reader would give them."""
    def read_zone(image, zone):
        return [w for line in lines for w in line
                if w.box.top < zone.bottom and w.box.bottom > zone.top]
    return read_zone


needs_font = pytest.mark.skipif(FONT is None, reason="no TrueType font")
ROWS = ["Relatório de Admissão", "Nome: Maria Santos", "NIF: 123456789",
        "Email: maria@exemplo.pt"]


def known_lines():
    """The words of document_image() with the exact boxes they were drawn
    in, so the test does not depend on what Tesseract reads."""
    font = ImageFont.truetype(str(FONT), 40)
    lines = []
    for row, text in enumerate(ROWS):
        words = text.split(" ")
        line = []
        for i, word in enumerate(words):
            prefix = " ".join(words[:i]) + (" " if i else "")
            x = 50 + font.getlength(prefix)
            left, top, right, bottom = font.getbbox(word)
            box = ocr.Box(round(x + left), 40 + row * 90 + top,
                          round(x + right), 40 + row * 90 + bottom)
            line.append(ocr.OcrWord(word, box, 95.0))
        lines.append(line)
    return lines


@needs_font
def test_ink_sticking_out_of_words_already_read_widens_them(monkeypatch):
    image = document_image()
    lines = known_lines()
    # Boxes too narrow for their words, as on a tilted line: the ink that
    # sticks out looks like an unread line, but the words are already there.
    narrow = [[ocr.OcrWord(w.text, ocr.Box(
        w.box.left, w.box.top,
        w.box.left + int((w.box.right - w.box.left) * 0.75),
        w.box.bottom), w.confidence) for w in line] for line in lines]
    assert ocr.unread_lines(image, lines) == 0
    assert ocr.unread_lines(image, narrow) >= 1
    monkeypatch.setattr(ocr, "_read_zone", reads_back(lines))
    merged, still_unread = ocr.recover_unread(image, narrow)
    assert still_unread == []
    assert first_words(merged) == first_words(narrow)
    assert [len(line) for line in merged] == [len(line) for line in narrow]
    # The words that stuck out now have the whole word in their box, so a
    # mask paints all of its ink; no box got smaller.
    before = [w.box for line in narrow for w in line]
    after = [w.box for line in merged for w in line]
    assert all(new.left <= old.left and new.right >= old.right
               for old, new in zip(before, after))
    assert any(new.right > old.right for old, new in zip(before, after))


def ink_bar():
    image = Image.new("RGB", (1500, 200), "white")
    ImageDraw.Draw(image).rectangle([50, 80, 900, 110], fill="black")
    return image


def test_a_zone_nothing_can_be_read_in_stays_unread(monkeypatch):
    monkeypatch.setattr(ocr, "_read_zone", lambda image, zone: [])
    merged, still_unread = ocr.recover_unread(ink_bar(), [])
    assert merged == []
    assert len(still_unread) == 1


def test_noise_without_letters_or_digits_is_not_a_recovered_line(monkeypatch):
    noise = [ocr.OcrWord("|", ocr.Box(60, 80, 70, 110), 90.0)]
    monkeypatch.setattr(ocr, "_read_zone", lambda image, zone: noise)
    merged, still_unread = ocr.recover_unread(ink_bar(), [])
    assert merged == []
    assert len(still_unread) == 1


def test_a_long_word_the_reader_is_unsure_of_is_still_recovered(monkeypatch):
    # An email read as "carlos.mendes(&centromedicalisboa.pt" at confidence
    # 36: the "@" confuses OCR, but the line is there and must be masked.
    email = [ocr.OcrWord("Email:", ocr.Box(60, 80, 160, 110), 90.0),
             ocr.OcrWord("carlos.mendes(&centromedicalisboa.pt",
                         ocr.Box(170, 80, 912, 110), 36.0)]
    monkeypatch.setattr(ocr, "_read_zone", lambda image, zone: email)
    merged, still_unread = ocr.recover_unread(ink_bar(), [])
    assert [w.text for w in merged[0]] == [w.text for w in email]
    assert still_unread == []


def test_a_short_unsure_word_is_noise(monkeypatch):
    weak = [ocr.OcrWord("a1b", ocr.Box(60, 80, 400, 110), 30.0)]
    monkeypatch.setattr(ocr, "_read_zone", lambda image, zone: weak)
    merged, still_unread = ocr.recover_unread(ink_bar(), [])
    assert merged == []
    assert len(still_unread) == 1


def test_low_confidence_words_are_not_recovered(monkeypatch):
    weak = [ocr.OcrWord("abc", ocr.Box(60, 80, 400, 110), 10.0)]
    monkeypatch.setattr(ocr, "_read_zone", lambda image, zone: weak)
    merged, still_unread = ocr.recover_unread(ink_bar(), [])
    assert merged == []
    assert len(still_unread) == 1


def test_short_numbers_are_not_uncertain_identifiers():
    from anonimizador.sources import looks_like_identifier
    assert not any(map(looks_like_identifier, ["72kg", "50mg", "BRCA1"]))
    assert all(map(looks_like_identifier,
                   ["123456/89", "mariaQgexemplo.pt", "x@y", "912"]))


def test_ocr_fixes_are_taken_word_by_word_only_when_safe():
    from anonimizador.ocr_fix import fix_lines, merge
    assert merge("Informações Financelras:", "Informações Financeiras:") \
        == "Informações Financeiras:"
    # A different word count would move the bars: keep the OCR reading.
    assert merge("Peso: 72Kkg", "Peso: 72 kg") == "Peso: 72Kkg"
    # A different name is not a typo fix.
    assert merge("Nome: Maria", "Nome: Joana") == "Nome: Maria"

    def fixer(prompt, **kwargs):
        return '{"lines": ["CVV 123", "Rendimento: €62.500"]}'

    assert fix_lines(["OVV 123", "Rendimento; €62,500"],
                     generate_fn=fixer) == ["CVV 123", "Rendimento: €62.500"]

    def broken(prompt, **kwargs):
        raise RuntimeError("model down")

    assert fix_lines(["OVV 123"], generate_fn=broken) == ["OVV 123"]


R = ('xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
     'relationships"')


def test_docx_email_links_and_field_codes_are_removed():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = Document()
    paragraph = document.add_paragraph("Contacto: ")
    r_id = document.part.relate_to("mailto:maria@exemplo.pt", RT.HYPERLINK,
                                   is_external=True)
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {W} {R} r:id="{r_id}"><w:r>'
        f'<w:t>maria@exemplo.pt</w:t></w:r></w:hyperlink>'))
    field = document.add_paragraph("Site: ")
    field._p.append(parse_xml(
        f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>'))
    field._p.append(parse_xml(
        f'<w:r {W}><w:instrText xml:space="preserve"> HYPERLINK '
        f'"mailto:maria@exemplo.pt" </w:instrText></w:r>'))
    field._p.append(parse_xml(
        f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>'))
    field._p.append(parse_xml(
        f'<w:r {W}><w:t>escrever</w:t></w:r>'))
    field._p.append(parse_xml(
        f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>'))
    buffer = io.BytesIO()
    document.save(buffer)

    source = load_source(buffer.getvalue(), ".docx")
    assert source.lines == ["Contacto: maria@exemplo.pt", "Site: escrever"]
    result = redact(source, secret_words(source.lines), "contacto.docx")
    xml = all_xml_text(result.data)
    assert "maria@exemplo.pt" not in xml and "maria%40exemplo" not in xml
    assert "Contacto: *" in load_source(result.data, ".docx").lines


def test_docx_text_boxes_are_refused_with_a_clear_reason():
    document = Document()
    box = document.add_paragraph("Caixa: ")
    box._p.append(parse_xml(
        f'<w:r {W}><w:pict><w:txbxContent><w:p><w:r><w:t>Maria Santos'
        f'</w:t></w:r></w:p></w:txbxContent></w:pict></w:r>'))
    buffer = io.BytesIO()
    document.save(buffer)
    source = load_source(buffer.getvalue(), ".docx")
    assert "Maria Santos" in source.lines
    with pytest.raises(RedactionError, match="caixas de texto"):
        redact(source, secret_words(source.lines), "caixa.docx")


def test_long_scans_are_refused_before_ocr(monkeypatch):
    from anonimizador import sources

    monkeypatch.setattr(sources, "MAX_OCR_PAGES", 2)
    doc = pymupdf.open()
    for _ in range(3):
        doc.new_page()
    with pytest.raises(sources.TooManyPages):
        load_source(doc.tobytes(), ".pdf")


@needs_ocr
def test_webp_photo_is_read_and_redacted():
    buffer = io.BytesIO()
    document_image().save(buffer, format="WEBP", quality=90)
    source = load_source(buffer.getvalue(), ".webp")
    assert source.kind == "image"
    assert any("Maria Santos" in line for line in source.lines)
    result = redact(source, secret_words(source.lines), "foto.webp")
    assert result.filename == "foto_censurado.png"
    lines = ocr.line_texts(ocr.read_image(
        Image.open(io.BytesIO(result.data)).convert("RGB")))
    assert not any("Maria" in line for line in lines)
